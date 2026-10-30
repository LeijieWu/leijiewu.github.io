"""Export the general academic CV from the same records used by Jekyll."""
from pathlib import Path
from datetime import date
import re
import shutil
import hashlib
import yaml
from bs4 import BeautifulSoup, NavigableString
from markdown_it import MarkdownIt
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'Full_CV.docx'
INK = '263238'
ACCENT = '238CA5'
MUTED = '607078'

def records(folder):
    out = []
    for path in (ROOT / folder).glob('*.md'):
        text = path.read_text(encoding='utf-8')
        data = yaml.safe_load(text.split('---', 2)[1])
        out.append(data)
    return sorted(out, key=lambda x: str(x['date']), reverse=True)

def cv_html():
    source = (ROOT / '_includes/cv-content.md').read_text(encoding='utf-8')
    source = source.replace('{% include base_path %}', '')
    pubs = records('_publications')
    for kind, var in [('journal', 'journal_publications'), ('conference', 'conference_publications'), ('preprint', 'preprint_publications')]:
        rows = []
        for item in pubs:
            if item['publication_type'] != kind:
                continue
            rows.append('<li class="publication">' + item['citation'] + ' <a href="' + item['paperurl'] + '">[Paper]</a></li>')
        pattern = r'<ul>\{% assign ' + var + r' = .*?\{% endfor %\}</ul>'
        source, count = re.subn(pattern, '<ul>' + ''.join(rows) + '</ul>', source, flags=re.S)
        assert count == 1, (kind, count)
    talks = []
    for item in records('_talks'):
        title = item['title'].removeprefix('Conference Talk: ')
        when = item['date'].strftime('%B %Y')
        talks.append('<li><b>' + title + '</b>. ' + item['type'] + ', ' + item['location'] + ', ' + when + '. Presented by ' + item['presenter'] + '. <a href="' + item['paperurl'] + '">[Paper]</a></li>')
    source, count = re.subn(r'<ul>\{% for post in site.talks reversed %\}.*?\{% endfor %\}</ul>', '<ul>' + ''.join(talks) + '</ul>', source, flags=re.S)
    assert count == 1
    assert '{%' not in source and '{{' not in source
    return MarkdownIt('commonmark', {'html': True}).render(source)

def field(paragraph, name):
    r = paragraph.add_run()
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), name)
    r._r.addnext(f)

def rich(paragraph, node, bold=False, italic=False, link=None):
    if isinstance(node, NavigableString):
        text = str(node)
        if not text or text.isspace():
            return
        text = re.sub(r'\s+', ' ', text)
        if link:
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            col = OxmlElement('w:color'); col.set(qn('w:val'), ACCENT); props.append(col)
            if bold: props.append(OxmlElement('w:b'))
            if italic: props.append(OxmlElement('w:i'))
            run.append(props); t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
            run.append(t); hyperlink.append(run); paragraph._p.append(hyperlink)
        else:
            run = paragraph.add_run(text); run.bold = bold; run.italic = italic
        return
    if node.name in ('ul', 'ol'):
        return
    if node.name == 'br':
        paragraph.add_run().add_break(); return
    for child in node.children:
        rich(paragraph, child, bold or node.name in ('b', 'strong'), italic or node.name in ('i', 'em'), node.get('href') if node.name == 'a' else link)

def export():
    if OUTPUT.exists():
        digest = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()[:12]
        archive = ROOT / 'local' / 'cv_archive' / ('Full_CV_' + digest + '.docx')
        archive.parent.mkdir(parents=True, exist_ok=True)
        if not archive.exists(): shutil.copy2(OUTPUT, archive)
    cfg = yaml.safe_load((ROOT / '_config.yml').read_text(encoding='utf-8'))
    doc = Document()
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(16); sec.bottom_margin = Mm(16)
    sec.left_margin = Mm(17); sec.right_margin = Mm(17)
    sec.header_distance = Mm(7); sec.footer_distance = Mm(8)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = Pt(12.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for key in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[key].font.name = 'Arial'
    title = doc.styles['Title']; title.font.size = Pt(25); title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(4)
    subtitle = doc.styles['Subtitle']; subtitle.font.size = Pt(10.5); subtitle.font.color.rgb = RGBColor.from_string(ACCENT)
    for key, size in [('Heading 1', 11), ('Heading 2', 10.3), ('Heading 3', 10.3)]:
        st = doc.styles[key]; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(ACCENT)
        st.paragraph_format.space_before = Pt(10); st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    doc.add_paragraph('Leijie Wu', 'Title')
    doc.add_paragraph('Research Scientist | Huawei Hong Kong Research Center', 'Subtitle')
    c = doc.add_paragraph()
    author = cfg['author']
    rich(c, BeautifulSoup('<a href="mailto:'+author['email']+'">'+author['email']+'</a> | '+author['phone']+' | Hong Kong SAR<br/><a href="'+cfg['url']+'">leijiewu.github.io</a> | <a href="'+author['googlescholar']+'">Google Scholar</a> | <a href="'+author['orcid']+'">ORCID</a>', 'html.parser'))
    for run in c.runs: run.font.size = Pt(9)
    soup = BeautifulSoup(cv_html(), 'html.parser')
    current = ''
    def block(node, level=0):
        nonlocal current
        if isinstance(node, NavigableString):
            if node.strip():
                para = doc.add_paragraph(); rich(para, node)
            return
        if node.name in ('h1','h2','h3'):
            text = node.get_text(' ', strip=True)
            if node.name == 'h1': current = text
            para = doc.add_paragraph(text, 'Heading 1' if node.name=='h1' else 'Heading 2')
            if text in ('Publications', 'Teaching Experience'):
                para.paragraph_format.page_break_before = True
            return
        if node.name in ('ul','ol'):
            for li in node.find_all('li', recursive=False): block(li, level)
            return
        if node.name in ('p','li','small'):
            para = doc.add_paragraph()
            is_pub = 'publication' in node.get('class', [])
            if node.name == 'li':
                para.paragraph_format.left_indent = Mm(3.5 + level*3)
                para.paragraph_format.first_line_indent = Mm(-3.5)
                para.add_run('- ')
            rich(para, node)
            para.paragraph_format.keep_together = True
            if is_pub:
                for run in para.runs: run.font.size = Pt(9)
                para.paragraph_format.line_spacing = Pt(11.5)
                para.paragraph_format.space_after = Pt(5)
            nested = node.find_all(['ul','ol'],recursive=False)
            if nested:
                para.paragraph_format.keep_with_next = True
                for child in nested: block(child, level+1)
            return
        for child in node.children: block(child,level)
    for node in soup.children: block(node)
    footer = sec.footer.paragraphs[0]
    footer.add_run('Leijie Wu | Full Academic CV | '+date.today().strftime('%B %Y')+'    ')
    field(footer,'PAGE'); footer.add_run(' / '); field(footer,'NUMPAGES')
    for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(MUTED)
    doc.core_properties.title='Leijie Wu Full Academic Curriculum Vitae'
    doc.core_properties.author='Leijie Wu'
    doc.core_properties.subject='Research, industrial technology transfer, publications, teaching and academic service'
    doc.save(OUTPUT)
    reopened=Document(OUTPUT)
    text=' '.join(p.text for p in reopened.paragraphs)
    assert 'Efficient AI systems' in text and 'Research Agenda' in text
    assert 'Present' in text
    assert 'Research Supervision and Mentoring' in text
    assert '{%' not in text and '{{' not in text
    print('Exported',OUTPUT,'paragraphs',len(reopened.paragraphs))

if __name__=='__main__':export()
